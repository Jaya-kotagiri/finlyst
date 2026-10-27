"""Read-only coverage audit of original sources and every completed pipeline stage."""

from collections import Counter
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from .classification import build
from .kotak import parse_kotak
from .sbi import parse_sbi
from .notes import blank, label, normalize_day, read_notes
from .reconcile import narration
from .output_models import (BankTransactions, ClassificationOutput, CoverageOutput,
                            EventDocument, ReconciliationOutput, RelationshipsOutput)
from .schema_validation import parse_document, validate_schema

ROOT = Path(__file__).resolve().parents[1]


def inspect_workbook(path):
    """Independently scan all rows, including beyond footers and hidden sheets."""
    workbook = load_workbook(path, data_only=False)
    transactions, other_notes = {}, []
    try:
        for sheet in workbook:
            columns = None
            for cells in sheet.iter_rows():
                labels = [label(c.value) for c in cells]
                if 'meaning' in labels and ('date' in labels or 'transaction date' in labels):
                    columns = {'date': labels.index('date' if 'date' in labels else 'transaction date'),
                               'meaning': labels.index('meaning'),
                               'follow_up': labels.index('missed') if 'missed' in labels else None}
                    continue
                if columns is None:
                    continue
                def cell(key):
                    index = columns[key]
                    return cells[index] if index is not None else None
                meaning, follow = cell('meaning'), cell('follow_up')
                record = {'sheet': sheet.title, 'row': cells[0].row,
                          'meaning_cell': meaning.coordinate, 'meaning': meaning.value,
                          'follow_up': follow.value if follow else None,
                          'hidden': sheet.sheet_state != 'visible' or bool(sheet.row_dimensions[cells[0].row].hidden)}
                try:
                    record['date'] = normalize_day(cell('date').value)
                except ValueError:
                    if not blank(record['meaning']) or not blank(record['follow_up']):
                        other_notes.append(record)
                    continue
                transactions[f'{sheet.title}:{cells[0].row}'] = record
    finally:
        workbook.close()
    return transactions, other_notes


def candidates(bank, notes):
    """Cross-check workbook presence independently of the saved 3A match."""
    return [n for n in notes if (n['date'], n['amount'], n['direction']) ==
            (bank['transaction_date'], bank['amount'], bank['direction']) and
            narration(n.get('narration')) == narration(bank['raw_description'])]


def reason(row):
    if row['effective_treatment'] != 'unknown':
        return 'supported', None, 'Existing evidence supports effective treatment.'
    if not row['workbook_candidates']:
        return 'missing_workbook_row', 'source_notes_coverage', 'No matching transaction row in supplied notes workbook; no semantic meaning available.'
    if len(row['workbook_candidates']) != 1:
        return 'ambiguous_workbook_match', 'source_alignment', 'Multiple workbook candidates; no forced association.'
    if not row['event_id']:
        return 'missing_generated_event', 'stage_2', 'Workbook transaction has no generated event.'
    if not row['matched_in_3A']:
        return 'unmatched_generated_event', 'stage_3A', 'Generated event is not attached to this bank transaction.'
    if not row['meaning_populated'] and not row['follow_up_populated']:
        return 'blank_note', 'source_notes_content', 'Narration/entity alone does not establish financial treatment; meaning cell is blank.'
    if row['classification_rule'] == 'funding_or_beneficiary_terms_unresolved':
        return 'uncertain_financial_terms', 'source_meaning_and_3B', 'Note retained; payer/beneficiary or funding terms do not establish expense, gift, loan or recoverable payment.'
    return 'semantic_classification_review', 'stage_3C1', 'Meaning exists but current classification rules do not support a treatment; inspect evidence before changing it.'


def coverage_counts(rows):
    unknown = {r['transaction_id'] for r in rows if r['effective_treatment'] == 'unknown'}
    unmatched = {r['transaction_id'] for r in rows if not r['matched_in_3A']}
    return {'transactions': len(rows), 'unknown': len(unknown), 'unmatched': len(unmatched),
            'unknown_and_unmatched': len(unknown & unmatched),
            'unknown_but_matched': len(unknown - unmatched),
            'unmatched_but_classified': len(unmatched - unknown),
            'original_unknown_ids': sorted(unknown),
            'unknown_with_populated_meaning': sum(r['meaning_populated'] for r in rows if r['transaction_id'] in unknown),
            'unknown_by_reason': dict(sorted(Counter(r['reason_category'] for r in rows if r['transaction_id'] in unknown).items())),
            'by_reason': dict(sorted(Counter(r['reason_category'] for r in rows).items()))}


def run(root=ROOT):
    root = Path(root)
    protected = sorted(p for folder in ('input', 'output/bank', 'output/events', 'output/reconciled',
                                        'output/relationships', 'output/classification')
                       for p in (root / folder).rglob('*') if p.is_file())
    hashes = {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in protected}
    def read(path, schema):
        return parse_document((root / path).read_bytes(), schema, str(root / path))
    documents = [read(f'output/reconciled/{b}_reconciled.json', ReconciliationOutput) for b in ('kotak', 'sbi')]
    view = read('output/classification/effective_transactions.json', ClassificationOutput)
    if build(documents, read('output/relationships/relationships.json', RelationshipsOutput)) != view:
        raise ValueError('Saved classification differs from validated deterministic rebuild')
    effective = {r['transaction_id']: r for r in view['transactions']}
    rows, inventory = [], {}
    for bank, parser, pattern in [('kotak', parse_kotak, '*.csv'), ('sbi', parse_sbi, 'sbi_unencrypted.*')]:
        paths = list((root / f'input/bank/{bank}').glob(pattern))
        if len(paths) != 1:
            raise ValueError(f'{bank}: source selection is not unique')
        source = paths[0]
        transactions = parser(source)
        if transactions != read(f'output/bank/{bank}_transactions.json', BankTransactions):
            raise ValueError('Original bank source differs from bank JSON')
        document = next(d for d in documents if d['bank'] == bank.upper())
        if [r['bank_transaction'] for r in document['transactions']] != transactions:
            raise ValueError('Bank facts differ in reconciled output')
        event_doc = read(f'output/events/{bank}_events.json', EventDocument)
        if document['event_source_document'] != event_doc:
            raise ValueError('Reconciled event source differs from final events')
        note_paths = list((root / f'input/notes/{bank}').glob('*.xlsx'))
        if len(note_paths) != 1:
            raise ValueError('Notes source selection is not unique')
        notes_path = note_paths[0]
        notes = read_notes(notes_path)
        raw, outside = inspect_workbook(notes_path)
        if set(raw) != {n['source_id'] for n in notes} or notes != event_doc['source_rows']:
            raise ValueError('Workbook/extracted/saved note coverage differs; investigate before reporting')
        for n in notes:
            original = raw[n['source_id']]
            if (n.get('note') or None) != (original['meaning'] or None) or (n.get('follow_up') or None) != (original['follow_up'] or None):
                raise ValueError('Original note text changed')
        events = {e['source_id']: e for e in event_doc['events']}
        if len(events) != len(event_doc['events']):
            raise ValueError('Duplicate event source IDs')
        # Original bank positions, independently cross-checked against parser counts.
        if bank == 'kotak':
            with source.open(encoding='utf-8-sig', newline='') as f:
                locations = [{'csv_record': i} for i, r in enumerate(csv.reader(f), 1)
                             if len(r) == 9 and r[0].isdigit()]
        else:
            workbook = load_workbook(source, read_only=True)
            locations = []
            try:
                for sheet in workbook:
                    for cells in sheet.iter_rows():
                        try:
                            normalize_day(cells[0].value)
                        except ValueError:
                            continue
                        locations.append({'sheet': sheet.title, 'row': cells[0].row})
            finally:
                workbook.close()
        if len(locations) != len(transactions):
            raise ValueError('Original bank row inventory differs')
        assigned_sources = []
        for transaction, location, reconciled in zip(transactions, locations, document['transactions']):
            identifier = transaction['transaction_id']
            possible = candidates(transaction, notes)
            note = possible[0] if len(possible) == 1 else None
            event = events.get(note['source_id']) if note else None
            attached = reconciled['event']
            if event != attached:
                raise ValueError(f'{identifier}: independent workbook association disagrees with 3A')
            if note:
                assigned_sources.append(note['source_id'])
            result = effective[identifier]
            row = dict(transaction, bank_source_file=str(source.relative_to(root)), bank_source_location=location,
                       notes_file=str(notes_path.relative_to(root)), workbook_candidates=[n['source_id'] for n in possible],
                       workbook_row=raw[note['source_id']] if note else None,
                       meaning_populated=bool(note and not blank(note.get('note'))),
                       follow_up_populated=bool(note and not blank(note.get('follow_up'))),
                       extracted_in_stage_2=note is not None, event_id=event['event_id'] if event else None,
                       matched_in_3A=attached is not None, semantic_meaning_preserved=event == attached,
                       effective_treatment=result['effective']['treatment'], classification_rule=result['effective']['rule'],
                       original_unknown=result['effective']['treatment'] == 'unknown',
                       semantic=result['semantic'], relationships=result['relationships'], review_reasons=result['review_reasons'])
            row['reason_category'], row['responsible_stage'], row['explanation'] = reason(row)
            rows.append(row)
        if len(assigned_sources) != len(set(assigned_sources)) or set(assigned_sources) != set(raw):
            raise ValueError('Workbook rows are not uniquely accounted for in bank transactions')
        inventory[bank] = {'bank_file': str(source.relative_to(root)), 'bank_count': len(transactions),
                           'bank_date_range': [min(t['transaction_date'] for t in transactions), max(t['transaction_date'] for t in transactions)],
                           'notes_file': str(notes_path.relative_to(root)), 'notes_count': len(notes),
                           'notes_date_range': [min(n['date'] for n in notes), max(n['date'] for n in notes)],
                           'populated_meanings': sum(not blank(n.get('note')) for n in notes),
                           'generated_events': len(events), 'unextracted_transaction_rows': 0,
                           'nontransaction_meaning_cells': outside,
                           'missing_workbook_ids': [r['transaction_id'] for r in rows if r['bank'] == bank.upper() and not r['workbook_candidates']]}
    if any(hashlib.sha256(p.read_bytes()).hexdigest() != hashes[str(p.relative_to(root))] for p in protected):
        raise ValueError('Protected source changed during audit')
    report = {'schema_version': '3C.1A.1', 'summary': coverage_counts(rows), 'sources': inventory,
              'verification': {'source_sha256': hashes, 'sources_unchanged': True,
                               'classification_rebuild_identical': True, 'deterministic_fixes': 0,
                               'api_calls': 0}, 'transactions': rows}
    validate_schema(CoverageOutput, report, 'coverage audit output')
    output = root / 'output/audit/coverage_audit.json'
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return report


if __name__ == '__main__':
    print(json.dumps(run()['summary'], indent=2))
