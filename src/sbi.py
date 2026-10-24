"""Convert the inspected SBI statement workbook to shared bank-fact JSON."""

import argparse
from datetime import date, datetime
import json
from pathlib import Path
import re
from zipfile import BadZipFile

from openpyxl import load_workbook

from .models import Transaction, money
from .validation import summary, validate_transaction, write_json

HEADER = ('Date', 'Details', 'Ref No/Cheque No', 'Debit', 'Credit', 'Balance')
SUMMARY_HEADER = ('Brought Forward (\u20b9)', 'Dr Count', 'Cr Count', 'Total Debits (\u20b9)',
                  'Total Credits (\u20b9)', 'Closing Balance (\u20b9)')
FOOTER = (
    'Please do not share your ATM, Debit/Credit Card number, PIN (Personal Identification number), OTP (One-Time Password), Username or Password with \nanyone via email, SMS, phone call, or any other medium. Bank never asks for such information.',
    'If your account is operated by a Power of Attorney holder, please review the transactions with extra care.',
    'This is a computer generated statement and does not require a signature.',
)


def blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def sbi_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return datetime.strptime(value.strip(), '%d/%m/%Y').date().isoformat()


def parse_sbi(path: Path) -> list[Transaction]:
    contents = Path(path).read_bytes()
    if contents.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
        raise ValueError('SBI workbook is encrypted or uses an unsupported OLE format; no JSON was written.')
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except BadZipFile as exc:
        raise ValueError('SBI input is not a readable XLSX workbook') from exc
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('SBI: expected one statement sheet; inspect additional sheets before conversion')
        sheet = workbook.worksheets[0]
        account = None
        transactions = []
        state = 'metadata'
        errors = []
        statement_summary = None
        for number, cells in enumerate(sheet.iter_rows(), 1):
            row = tuple(c.value for c in cells)
            if all(blank(v) for v in row):
                continue
            try:
                if any(c.data_type in ('f', 'e') for c in cells):
                    raise ValueError('formula or Excel error cell requires review')
                if len(row) > 6 and any(not blank(v) for v in row[6:]):
                    raise ValueError('unexpected populated columns')
                row = (row + (None,) * 6)[:6]
                if row == HEADER:
                    if state not in ('metadata', 'transactions'):
                        raise ValueError('transaction header after summary')
                    state = 'transactions'
                    continue
                if state == 'metadata':
                    for value in row:
                        if isinstance(value, str):
                            match = re.fullmatch(r'Account Number\s*:\s*(\d+)\s*', value)
                            if match:
                                if account is not None and account != match[1]:
                                    raise ValueError('conflicting account metadata')
                                account = match[1]
                    continue
                if state == 'transactions':
                    if isinstance(row[0], str) and re.fullmatch(
                        r'Statement Summary\s*:\s*\d{2}-\d{2}-\d{4}\s+To\s+\d{2}-\d{2}-\d{4}', row[0]
                    ) and all(blank(v) for v in row[1:]):
                        state = 'summary_header'
                        continue
                    debit, credit = not blank(row[3]), not blank(row[4])
                    if debit == credit:
                        raise ValueError('expected exactly one populated debit or credit amount')
                    if not isinstance(row[1], str):
                        raise ValueError('narration must be source text')
                    if not blank(row[2]) and not isinstance(row[2], str):
                        raise ValueError('non-text reference requires review')
                    transaction = Transaction(
                        transaction_id=f'SBI_{len(transactions) + 1:04d}', bank='SBI',
                        account=account, transaction_date=sbi_date(row[0]),
                        amount=money(row[3] if debit else row[4]),
                        direction='debit' if debit else 'credit', raw_description=row[1],
                        reference=None if blank(row[2]) else row[2],
                        balance_after=None if blank(row[5]) else money(row[5], signed=True),
                    )
                    validate_transaction(transaction)
                    transactions.append(transaction)
                elif state == 'summary_header':
                    if row != SUMMARY_HEADER:
                        raise ValueError('unrecognized summary header')
                    state = 'summary_values'
                elif state == 'summary_values':
                    for value in (row[0], row[5]):
                        if not isinstance(value, str) or not value.endswith(('CR', 'DR')):
                            raise ValueError('unrecognized summary balance')
                        money(value[:-2])
                    statement_summary = (int(row[1]), int(row[2]), money(row[3]), money(row[4]))
                    state = 'footer'
                elif state == 'footer':
                    if row[0] not in FOOTER or any(not blank(v) for v in row[1:]):
                        raise ValueError('unrecognized footer row')
            except (ValueError, TypeError, AttributeError) as exc:
                errors.append(f'SBI sheet {sheet.title!r} row {number}: {exc}')
        if errors:
            raise ValueError('Unsafe rows; no JSON written:\n' + '\n'.join(errors))
        if not transactions or state != 'footer' or statement_summary is None:
            raise ValueError('SBI: incomplete statement header, transactions, or summary')
        report = summary(transactions)
        actual = (report['debit']['count'], report['credit']['count'],
                  report['debit']['total'], report['credit']['total'])
        if actual != statement_summary:
            raise ValueError('SBI: parsed counts/totals differ from the source statement summary')
        return transactions
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parents[1] / 'output/bank/sbi_transactions.json')
    args = parser.parse_args()
    try:
        print(json.dumps(write_json(parse_sbi(args.source), args.output), indent=2))
    except ValueError as exc:
        parser.exit(1, f'{exc}\n')


if __name__ == '__main__':
    main()
