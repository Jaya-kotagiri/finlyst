"""Read annotated XLSX tables once; normalize facts without interpreting notes."""

from datetime import date, datetime
from pathlib import Path
import re

from openpyxl import load_workbook

from .models import money
from .kotak import FOOTER as KOTAK_FOOTER
from .sbi import FOOTER as SBI_FOOTER


def blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def label(value):
    return str(value or '').strip().lower().rstrip('.')


def normalize_day(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    for fmt in ('%d/%m/%Y', '%d-%m-%Y %H:%M:%S', '%d-%m-%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError('unrecognized transaction date')


def read_notes(path: Path) -> list[dict]:
    """Detect columns by labels, not bank, sheet name or fixed row offsets."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    result = []
    try:
        for sheet in workbook:
            columns = None
            footer = False
            summary_remaining = 0
            for number, cells in enumerate(sheet.iter_rows(), 1):
                row = [c.value for c in cells]
                if all(blank(v) for v in row):
                    continue
                labels = [label(v) for v in row]
                if 'meaning' in labels and ('date' in labels or 'transaction date' in labels):
                    if columns is not None:
                        raise ValueError(f'{sheet.title} row {number}: repeated notes header')
                    columns = {}
                    for key, aliases in {
                        'date': ('transaction date', 'date'), 'narration': ('description', 'details'),
                        'note': ('meaning',), 'follow_up': ('missed',), 'amount': ('amount',),
                        'direction': ('dr / cr',), 'debit': ('debit',), 'credit': ('credit',),
                        'reference': ('chq /ref no', 'ref no/cheque no'), 'value_date': ('value date',),
                    }.items():
                        columns[key] = next((labels.index(a) for a in aliases if a in labels), None)
                    if columns['narration'] is None:
                        raise ValueError('notes table has no narration column')
                    continue
                if columns is None:
                    continue
                try:
                    if any(c.data_type in ('f', 'e') for c in cells):
                        raise ValueError('formula/error cell needs review')
                    populated = [v for v in row if not blank(v)]
                    first = str(populated[0])
                    if first == 'Closing Balance':
                        footer = True
                        continue
                    if first.startswith('Statement Summary :'):
                        footer = True
                        summary_remaining = 2
                        continue
                    if footer:
                        if summary_remaining:
                            summary_remaining -= 1
                            continue
                        if len(populated) == 1 and first in (*KOTAK_FOOTER, *SBI_FOOTER):
                            continue
                        raise ValueError('unrecognized footer; row not discarded')
                    def get(key):
                        index = columns[key]
                        return row[index] if index is not None else None
                    item = {'source_id': f'{sheet.title}:{number}', 'date': normalize_day(get('date'))}
                    if columns['amount'] is not None:
                        item['amount'] = money(get('amount'))
                        marker = str(get('direction')).strip().upper()
                        if marker not in ('DR', 'CR'):
                            raise ValueError('unknown debit/credit marker')
                        item['direction'] = {'DR': 'debit', 'CR': 'credit'}[marker]
                    else:
                        debit, credit = get('debit'), get('credit')
                        if blank(debit) == blank(credit):
                            raise ValueError('expected one debit or credit')
                        item['amount'] = money(credit if blank(debit) else debit)
                        item['direction'] = 'credit' if blank(debit) else 'debit'
                    for key in ('narration', 'note', 'follow_up', 'reference'):
                        value = get(key)
                        if not blank(value):
                            if key == 'reference' and type(value) in (int, float) and float(value).is_integer():
                                value = str(int(value))
                                fmt = cells[columns[key]].number_format
                                if re.fullmatch(r'0+', fmt):
                                    value = value.zfill(len(fmt))
                            if not isinstance(value, str):
                                raise ValueError(f'{key} must be text')
                            item[key] = value
                    if not blank(get('value_date')):
                        item['value_date'] = normalize_day(get('value_date'))
                    result.append(item)
                except ValueError as exc:
                    raise ValueError(f'{sheet.title} row {number}: {exc}') from exc
            if columns is None and sheet.max_row > 1:
                raise ValueError(f'{sheet.title}: no recognized notes header')
        if not result:
            raise ValueError('no notes transaction rows')
        return result
    finally:
        workbook.close()
